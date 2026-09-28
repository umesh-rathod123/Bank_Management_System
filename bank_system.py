from openpyxl import load_workbook, Workbook
import os

FILE_NAME = "banking_system.xlsx"


# --------------------------------
# Excel File Setup
# --------------------------------
def setup_excel():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bank Accounts"

        sheet.append([
            "Account Number",
            "Name",
            "Mobile Number",
            "PIN",
            "Balance"
        ])

        workbook.save(FILE_NAME)


# --------------------------------
# Find Account
# --------------------------------
def find_account(account_number):
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row=row, column=1).value) == str(account_number):
            workbook.close()
            return row

    workbook.close()
    return None


# --------------------------------
# Verify PIN
# --------------------------------
def verify_pin(row, pin):
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    saved_pin = str(sheet.cell(row=row, column=4).value)

    workbook.close()

    return saved_pin == str(pin)


# --------------------------------
# Create Account
# --------------------------------
def create_account():
    print("\n========== CREATE ACCOUNT ==========")

    name = input("Enter Account Holder Name: ")
    mobile = input("Enter Mobile Number: ")

    while True:
        pin = input("Enter 4-digit PIN: ")

        if pin.isdigit() and len(pin) == 4:
            break
        else:
            print("Invalid PIN! Please enter exactly 4 digits.")

    while True:
        try:
            balance = float(input("Enter Initial Deposit: "))

            if balance >= 0:
                break
            else:
                print("Amount cannot be negative.")

        except ValueError:
            print("Please enter a valid amount.")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    if sheet.max_row == 1:
        account_number = 1001
    else:
        last_account = sheet.cell(
            row=sheet.max_row,
            column=1
        ).value

        account_number = int(last_account) + 1

    sheet.append([
        account_number,
        name,
        mobile,
        pin,
        balance
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nAccount created successfully!")
    print("Account Number:", account_number)
    print("Account Holder:", name)
    print("Initial Balance:", balance)


# --------------------------------
# Deposit Money
# --------------------------------
def deposit_money():
    print("\n========== DEPOSIT MONEY ==========")

    account_number = input("Enter Account Number: ")

    row = find_account(account_number)

    if row is None:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN!")
        return

    while True:
        try:
            amount = float(input("Enter Deposit Amount: "))

            if amount > 0:
                break
            else:
                print("Amount must be greater than 0.")

        except ValueError:
            print("Please enter a valid amount.")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    old_balance = float(sheet.cell(row=row, column=5).value)
    new_balance = old_balance + amount

    sheet.cell(row=row, column=5).value = new_balance

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nDeposit successful!")
    print("Updated Balance:", new_balance)


# --------------------------------
# Withdraw Money
# --------------------------------
def withdraw_money():
    print("\n========== WITHDRAW MONEY ==========")

    account_number = input("Enter Account Number: ")

    row = find_account(account_number)

    if row is None:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN!")
        return

    while True:
        try:
            amount = float(input("Enter Withdrawal Amount: "))

            if amount > 0:
                break
            else:
                print("Amount must be greater than 0.")

        except ValueError:
            print("Please enter a valid amount.")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    current_balance = float(sheet.cell(row=row, column=5).value)

    if amount > current_balance:
        print("\nInsufficient balance!")
        print("Available Balance:", current_balance)

        workbook.close()
        return

    new_balance = current_balance - amount

    sheet.cell(row=row, column=5).value = new_balance

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nWithdrawal successful!")
    print("Updated Balance:", new_balance)


# --------------------------------
# Check Balance
# --------------------------------
def check_balance():
    print("\n========== CHECK BALANCE ==========")

    account_number = input("Enter Account Number: ")

    row = find_account(account_number)

    if row is None:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN!")
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    name = sheet.cell(row=row, column=2).value
    balance = sheet.cell(row=row, column=5).value

    workbook.close()

    print("\nAccount Number:", account_number)
    print("Account Holder:", name)
    print("Current Balance:", balance)


# --------------------------------
# Main Menu
# --------------------------------
def main():

    setup_excel()

    while True:

        print("\n========================================")
        print("       SIMPLE BANKING SYSTEM")
        print("========================================")

        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("\nEnter your choice: ")

        if choice == "1":
            create_account()

        elif choice == "2":
            deposit_money()

        elif choice == "3":
            withdraw_money()

        elif choice == "4":
            check_balance()

        elif choice == "5":
            print("\nThank you for using Simple Banking System!")
            print("Have a nice day!")
            break

        else:
            print("\nInvalid choice! Please select 1 to 5.")


# --------------------------------
# Start Program
# --------------------------------
if __name__ == "__main__":
    main()